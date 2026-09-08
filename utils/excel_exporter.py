import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any

def export_applicants_to_excel(applicants: List[Dict[str, Any]], filename: str = "nomzodlar.xlsx") -> str:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Nomzodlar Ro'yxati"

    # Set gridlines visible explicitly
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "№ ID",
        "Sana va Vaqt",
        "F.I.O",
        "Telefon",
        "Telegram",
        "Yo'nalish (Vakansiya)",
        "Dasturlar va Darajalari",
        "Ish Tajribasi",
        "Ma'lumoti",
        "Yoshi",
        "Yashash Hududi",
        "Yashash Sharoiti",
        "Rus tili",
        "Ingliz tili",
        "Shaxsiy Kompyuter",
        "Haydovchilik",
        "Komandirovka",
        "Qo'shimcha Ishlash",
        "Kutilayotgan Maosh",
        "Portfolio / Fayl",
        "Holat (Status)"
    ]

    # Shriftlar va ranglar
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # To'q ko'k (Royal Blue)
    
    # Aniq va chiroyli chegaralar (Granitsalar)
    border_color = "94A3B8"  # Slate kulrang chegara
    cell_border = Border(
        left=Side(style='thin', color=border_color),
        right=Side(style='thin', color=border_color),
        top=Side(style='thin', color=border_color),
        bottom=Side(style='thin', color=border_color)
    )
    
    header_border = Border(
        left=Side(style='thin', color="0F2356"),
        right=Side(style='thin', color="0F2356"),
        top=Side(style='thin', color="0F2356"),
        bottom=Side(style='medium', color="0F2356")
    )

    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 1. Header qatorini bezash
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = header_border
    ws.row_dimensions[1].height = 32

    # Status ranglari
    status_styles = {
        "new": {
            "fill": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
            "font": Font(name="Segoe UI", size=10, bold=True, color="92400E"),
            "text": "🟡 Yangi"
        },
        "accepted": {
            "fill": PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid"),
            "font": Font(name="Segoe UI", size=10, bold=True, color="065F46"),
            "text": "🟢 Suhbatga taklif"
        },
        "rejected": {
            "fill": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
            "font": Font(name="Segoe UI", size=10, bold=True, color="991B1B"),
            "text": "🔴 Rad etildi"
        }
    }

    # 2. Ma'lumotlarni yozish
    row_idx = 2
    for app in applicants:
        username = f"@{app.get('username')}" if app.get('username') else "Mavjud emas"
        
        # Dasturlarni yangi qatorlar bilan formatlash
        raw_programs = app.get("programs_skill", "")
        if raw_programs and ", " in raw_programs and "\n" not in raw_programs:
            formatted_programs = "\n".join([f"• {p.strip()}" for p in raw_programs.split(",")])
        elif raw_programs and not raw_programs.startswith("•"):
            formatted_programs = f"• {raw_programs}"
        else:
            formatted_programs = raw_programs

        status_key = app.get("status", "new")
        status_info = status_styles.get(status_key, status_styles["new"])

        row_values = [
            f"#{app.get('id')}",
            str(app.get("created_at", "")),
            app.get("full_name", ""),
            app.get("phone", ""),
            username,
            app.get("direction", ""),
            formatted_programs,
            app.get("experience", ""),
            app.get("education", ""),
            app.get("age_range", ""),
            app.get("region", ""),
            app.get("housing", ""),
            app.get("russian_level", ""),
            app.get("english_level", ""),
            app.get("device", ""),
            app.get("driving", ""),
            app.get("trip_ready", ""),
            app.get("overtime_ready", ""),
            app.get("expected_salary", ""),
            app.get("portfolio", ""),
            status_info["text"]
        ]

        # Qatorlar uchun yumshoq zebra rangi
        is_even = (row_idx % 2 == 0)
        default_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid") if is_even else PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        # Qator balandligini dasturlar soniga qarab dinamik belgilash
        num_lines = max(len(str(val).split("\n")) for val in row_values)
        calc_height = max(26, num_lines * 18)

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = cell_border
            
            # Status ustuni (21-ustun) uchun maxsus rang
            if col_idx == 21:
                cell.fill = status_info["fill"]
                cell.font = status_info["font"]
                cell.alignment = center_align
            else:
                cell.fill = default_fill
                cell.font = Font(name="Segoe UI", size=10)
                # Markazlashtiriladigan ustunlar
                if col_idx in [1, 2, 4, 5, 10, 12, 13, 14, 17, 18]:
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

        ws.row_dimensions[row_idx].height = calc_height
        row_idx += 1

    # 3. Ustun kengliklarini avtomatik to'g'irlash
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            lines = val_str.split('\n')
            max_line = max(len(l) for l in lines) if lines else len(val_str)
            if max_line > max_len:
                max_len = max_line
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 13), 50)

    # 4. Header qatorini muzlatish (skroll qilganda ko'rinib turishi uchun)
    ws.freeze_panes = "A2"

    filepath = os.path.abspath(filename)
    wb.save(filepath)
    return filepath
